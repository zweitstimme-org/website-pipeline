#!/usr/bin/env python3
"""Live Berlin AGH 2026 nowcast from AfS _A_/_W_ CSVs.

Local history: AGH 2023 remapped to 2026 WKs (panel) blended with BTW 2025
Wahlbezirke when the AfS remap workbook is present. π₀ = forecast_state_be.
Falls back to forecast-only if live CSVs are still empty.
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path

from parliament_size_sim import allocate_be
from wahlabend_live_common import (
    AFS_2026_PXX,
    BE_BEZIRKSLISTE_PARTIES,
    BEZ_NAMES,
    EPS,
    EXTERNAL_JSON,
    MAIN,
    NAME_TO_BEZ,
    OUT_DIR,
    PARTIES,
    PARTY_LABELS,
    REPO,
    _load_json,
    _load_json_optional,
    _num,
    _pct,
    _shares,
    awk_code,
    blend_wk_history,
    blend_with_external,
    clock_from_fields,
    counts_from_row,
    district_priors,
    district_wk_unc_pp,
    github_now,
    live_meta,
    live_precincts,
    load_awk_to_wkr,
    load_direkt_roster,
    load_external,
    load_listen_roster,
    load_state_draws,
    load_wkr_panel,
    map_party_columns,
    merge_history,
    night_entry_mc,
    union_history_payloads,
    night_scenario_probs,
    nowcast_wkrs,
    precinct_ist_soll,
    prior_uncertainty_pp,
    read_csv_rows,
    row_has_votes,
    statewide_prior,
    turnout_mixture,
    unc_phase,
    wkr_races,
)

LIVE_DIR = REPO / "berlin" / "wahlabend" / "live"
RAW_WB = REPO / "berlin" / "wahlabend" / "raw"
PROCESSED = REPO / "berlin" / "wahlabend" / "processed"
PANEL_PATH = REPO / "berlin" / "agh23_be_abs.csv"
HIST_JSON = PROCESSED / "hist_wk_btw2025.json"

FORECAST_STATE_URL = "https://zweitstimme.org/data/forecast_state_be.json"
FORECAST_DISTRICT_URL = "https://zweitstimme.org/data/forecast_districts_be.json"
FORECAST_DRAWS_URL = "https://zweitstimme.org/data/forecast_state_be_draws.json"

AFS_BASES = (
    "https://www.wahlen-berlin.de/wahlen/Be2026/AFSPRAES/agh",
    "https://www.wahlen-berlin.de/wahlen/BE2026/AFSPRAES/agh",
)

LAST_OFFICE_TURNOUT = 62.9  # AGH 2023
RECENT_FEDERAL_TURNOUT = 78.0  # Berlin BTW 2025 (approx.)
TURNOUT_PRIOR = 70.0
WBER_2026 = 2_450_000
# AfS land-row AnzWbez on the 2026 _A_ testdata (78 WK × ~50 WB).
EXPECTED_PRECINCTS = 4114

KIND_LABEL = {
    "empty": "Noch keine Auszählung",
    "Z": "Zwischenergebnis",
    "V": "Vorläufiges Ergebnis",
    "E": "Endgültiges Ergebnis",
    "A": "AfS-Abzug",
}

# AfS night feed is still _A_ only (_W_ stays 404). The HTML Präsentation
# does publish an Ankunftstafel of arriving Wahlbezirke (names + WK, no votes).
BE_WB_NIGHT_NOTE = (
    "AfS liefert nachts keine Stimmen je Wahlbezirk (_W_ bleibt 404). "
    "Die Präsentation zeigt aber die Ankunftstafel — welche Wahlbezirke "
    "gerade eingegangen sind — plus Ist/Soll je Land, Bezirk und Wahlkreis. "
    "Die Suchseite «Ergebnisse nach Wahlbezirk» öffnet AfS erst nach "
    "vollständigem Eingang."
)
AFS_INDEX = "https://wahlen-berlin.de/wahlen/BE2026/Afspraes/agh/index.html"
ANKUNFT_ROW_RE = re.compile(
    r'data-sort="((?:\d{5}|\d{3}[A-Z]))\s+-\s+([^"]+)"'
    r'[\s\S]{0,800}?'
    r'ergebnisse_wahlkreis_(\d{4})\.html'
    r'[\s\S]{0,400}?'
    r'data-sort="(\d{1,2}:\d{2})"',
    re.I,
)


def ankunft_to_addr(code: str) -> str:
    """Präsentation IDs → AfS Adresse (12W404 / 08B5E)."""
    s = str(code or "").strip().upper()
    if re.fullmatch(r"\d{3}[A-Z]", s):
        return f"{s[:2]}B{s[2]}{s[3]}"
    if re.fullmatch(r"\d{5}", s):
        return f"{s[:2]}W{s[2]}{s[3:]}"
    return ""


def _html_text(s: str) -> str:
    return (
        str(s or "")
        .replace("&auml;", "ä")
        .replace("&ouml;", "ö")
        .replace("&uuml;", "ü")
        .replace("&Auml;", "Ä")
        .replace("&Ouml;", "Ö")
        .replace("&Uuml;", "Ü")
        .replace("&szlig;", "ß")
        .replace("&amp;", "&")
        .replace("&nbsp;", " ")
        .replace("&quot;", '"')
        .strip()
    )


def parse_ankunftstafel(html: str) -> list[dict]:
    """Last arrivals from the AfS Präsentation Ankunftstafel (no vote counts)."""
    awk_to_wkr = load_awk_to_wkr()
    out: list[dict] = []
    seen: set[str] = set()
    for m in ANKUNFT_ROW_RE.finditer(html or ""):
        code = m.group(1).strip().upper()
        if code in seen:
            continue
        seen.add(code)
        awk = m.group(3).strip()
        out.append(
            {
                "id": code,
                "name": _html_text(m.group(2)),
                "awk": awk,
                "wkr": awk_to_wkr.get(awk),
                "addr": ankunft_to_addr(code),
                "time": m.group(4).strip(),
                "art": "B" if re.fullmatch(r"\d{3}[A-Z]", code) else "W",
            }
        )
    return out


def merge_ankunft(path: Path, arrivals: list[dict]) -> dict:
    """Accumulate Ankunftstafel IDs across polls (the HTML only keeps ~10)."""
    doc: dict = {}
    if path.exists():
        try:
            doc = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            doc = {}
    stations = dict(doc.get("stations") or {})
    for rec in arrivals:
        key = str(rec.get("id") or "")
        if not key:
            continue
        prev = stations.get(key) or {}
        stations[key] = {
            **prev,
            **rec,
            "first_seen": prev.get("first_seen") or rec.get("time"),
        }
    out = {
        "stations": stations,
        "latest": arrivals,
        "n": len(stations),
        "source": AFS_INDEX,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(out, ensure_ascii=False), encoding="utf-8")
    return out


def load_ankunft(live_dir: Path | None = None) -> dict:
    html_path = (live_dir or LIVE_DIR) / "afs_index.html"
    json_path = (live_dir or LIVE_DIR) / "ankunft.json"
    arrivals: list[dict] = []
    if html_path.exists():
        arrivals = parse_ankunftstafel(
            html_path.read_text(encoding="utf-8", errors="replace")
        )
    return merge_ankunft(json_path, arrivals)


def _bezirk_of_name(name: str) -> str | None:
    s = str(name or "").strip().lower()
    for label, code in NAME_TO_BEZ.items():
        if s.startswith(label):
            return code
    return None


def attach_bezirke(panel: dict[str, dict]) -> None:
    awk = load_awk_to_wkr()
    wkr_to_bez = {wid: key[:2] for key, wid in awk.items()}
    for wid, row in panel.items():
        bez = wkr_to_bez.get(wid) or _bezirk_of_name(row.get("name") or "")
        row["bezirk"] = bez


def _xlsx_header_index(header: list[str]) -> dict[str, int]:
    return {re.sub(r"\s+", " ", h).strip().lower(): i for i, h in enumerate(header) if h}


def _xlsx_col(idx: dict[str, int], *names: str) -> int | None:
    for n in names:
        if n.lower() in idx:
            return idx[n.lower()]
    for key, i in idx.items():
        if any(n.lower() in key for n in names):
            return i
    return None


def _xlsx_rows(path: Path, sheet_hint: str | None = None) -> tuple[list, list[str]]:
    try:
        import openpyxl
    except ImportError:
        return [], []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return [], []
    sheet = None
    for name in wb.sheetnames:
        low = name.lower()
        if low.startswith("impress") or "erläut" in low:
            continue
        if sheet_hint and sheet_hint.lower() in low:
            sheet = name
            break
        if sheet is None:
            sheet = name
        if "zweit" in low or "w2" in low:
            sheet = name
            if not sheet_hint:
                break
    if sheet is None:
        wb.close()
        return [], []
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []
    header = [str(c).replace("\n", " ") if c is not None else "" for c in rows[0]]
    return rows, header


def _wid_from_bez_wk(bez, local, awk_to_wkr: dict[str, str]) -> str | None:
    key = awk_code(bez, local)
    if not key:
        return None
    return awk_to_wkr.get(key)


def _addr_key(val) -> str:
    if val is None:
        return ""
    return str(val).strip().split()[0]


def build_btw2025_history(panel: dict[str, dict]) -> tuple[dict[str, dict[str, float]], dict[str, str]]:
    """BTW 2025 on 2026 WBs → WK shares + Adresse→WK map.

    The official remap workbook folds BSW into Sonstige. BSW is taken from
    DL_BE_BU2025.xlsx (BE_W2) where 2025 Adressen overlap 2026 remapped WBs.
    """
    awk_to_wkr = load_awk_to_wkr()
    addr_to_wkr: dict[str, str] = {}
    buckets: dict[str, dict[str, float]] = {}

    remap = RAW_WB / "DL_BE_AGH2026_BT2025.xlsx"
    rows, header = _xlsx_rows(remap, "2025_zweit")
    if rows:
        idx = _xlsx_header_index(header)
        i_addr = _xlsx_col(idx, "adresse")
        i_bez = _xlsx_col(idx, "bezirksnummer")
        i_wk = _xlsx_col(idx, "abgeordnetenhauswahlkreis", "abgeordneten-hauswahlkreis")
        i_g = _xlsx_col(idx, "gültige stimmen")
        party_i = {
            "cdu": _xlsx_col(idx, "cdu"),
            "spd": _xlsx_col(idx, "spd"),
            "gruene": _xlsx_col(idx, "grüne", "gruene"),
            "linke": _xlsx_col(idx, "linke"),
            "afd": _xlsx_col(idx, "afd"),
            "fdp": _xlsx_col(idx, "fdp"),
            "bsw": _xlsx_col(idx, "bsw"),
        }
        for r in rows[1:]:
            if not r:
                continue
            bez = r[i_bez] if i_bez is not None else None
            loc = r[i_wk] if i_wk is not None else None
            wid = _wid_from_bez_wk(bez, loc, awk_to_wkr)
            if not wid:
                continue
            addr = _addr_key(r[i_addr]) if i_addr is not None else ""
            if addr:
                addr_to_wkr[addr] = wid
            slot = buckets.setdefault(wid, {p: 0.0 for p in PARTIES} | {"gueltig": 0.0})
            g = float(r[i_g] or 0) if i_g is not None else 0.0
            slot["gueltig"] += g
            named = 0.0
            for p, ci in party_i.items():
                if ci is None or ci >= len(r):
                    continue
                v = float(r[ci] or 0)
                slot[p] += v
                named += v
            slot["others"] += max(0.0, g - named)

    # Strukturdaten fills Adresse→WK for 2026 Urne WBs missing from the remap.
    struktur = RAW_WB / "DL_BE_AH2026_Strukturdaten.xlsx"
    srows, sheader = _xlsx_rows(struktur, "struktur")
    if srows and awk_to_wkr:
        idx = _xlsx_header_index(sheader)
        i_addr = _xlsx_col(idx, "adresse")
        i_bez = _xlsx_col(idx, "bezirksnummer")
        i_wk = _xlsx_col(idx, "abgeordnetenhauswahlkreis")
        for r in srows[1:]:
            if not r or i_addr is None or r[i_addr] is None:
                continue
            addr = _addr_key(r[i_addr])
            if not addr or addr in addr_to_wkr:
                continue
            wid = _wid_from_bez_wk(
                r[i_bez] if i_bez is not None else None,
                r[i_wk] if i_wk is not None else None,
                awk_to_wkr,
            )
            if wid:
                addr_to_wkr[addr] = wid

    # BSW from original BTW 2025 WB file (remap has no BSW column).
    bsw_path = RAW_WB / "DL_BE_BU2025.xlsx"
    brows, bheader = _xlsx_rows(bsw_path, "be_w2")
    n_bsw = 0
    if brows:
        idx = _xlsx_header_index(bheader)
        i_addr = _xlsx_col(idx, "adresse")
        i_bsw = _xlsx_col(idx, "bsw")
        i_bez = _xlsx_col(idx, "bezirksnummer")
        i_wk = _xlsx_col(idx, "abgeordneten- hauswahlkreis", "abgeordnetenhauswahlkreis")
        if i_bsw is not None:
            for r in brows[1:]:
                if not r:
                    continue
                v = float(r[i_bsw] or 0)
                if v <= 0:
                    continue
                addr = _addr_key(r[i_addr]) if i_addr is not None else ""
                wid = addr_to_wkr.get(addr)
                if not wid:
                    wid = _wid_from_bez_wk(
                        r[i_bez] if i_bez is not None else None,
                        r[i_wk] if i_wk is not None else None,
                        awk_to_wkr,
                    )
                if not wid or wid not in buckets:
                    continue
                # Move BSW out of others (remap lumped it into Sonstige).
                buckets[wid]["bsw"] += v
                buckets[wid]["others"] = max(0.0, buckets[wid]["others"] - v)
                n_bsw += 1

    out: dict[str, dict[str, float]] = {}
    for wid, slot in buckets.items():
        g = slot.pop("gueltig")
        if g <= 0:
            continue
        out[wid] = _shares(slot)

    if out:
        PROCESSED.mkdir(parents=True, exist_ok=True)
        HIST_JSON.write_text(
            json.dumps(
                {
                    "wkr": out,
                    "adresse_to_wkr": addr_to_wkr,
                    "n_wkr": len(out),
                    "n_wb": len(addr_to_wkr),
                    "n_bsw_rows": n_bsw,
                    "source": "DL_BE_AGH2026_BT2025.xlsx + DL_BE_BU2025.xlsx BE_W2",
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
    return out, addr_to_wkr


def load_btw2025_wk_shares(panel: dict[str, dict]) -> tuple[dict[str, dict[str, float]], dict[str, str]]:
    """Load cached 2026-WK BTW 2025 shares, rebuilding from xlsx if needed."""
    if HIST_JSON.exists():
        try:
            doc = json.loads(HIST_JSON.read_text(encoding="utf-8"))
            raw = doc.get("wkr") or {}
            out = {str(k): _shares(v) for k, v in raw.items() if isinstance(v, dict) and str(k) in panel}
            addr = {str(a): str(w) for a, w in (doc.get("adresse_to_wkr") or {}).items()}
            if len(out) >= 70:
                return out, addr
        except json.JSONDecodeError:
            pass
    return build_btw2025_history(panel)


def _looks_empty_afs(rows: list[dict]) -> bool:
    if not rows:
        return True
    return not any(row_has_votes(row) for row in rows)


def _afs_clock(row: dict) -> str | None:
    return clock_from_fields(row.get("Datum") or "", row.get("Zeit") or row.get("Uhrzeit") or "")


def _afs_unit(row: dict, party_cols: dict[str, str]) -> dict:
    counts, gueltig = counts_from_row(
        row,
        party_cols,
        ["Gueltig", "Gültige Stimmen", "Gültige", "Gültig"],
    )
    wber = _num(row.get("WberIns") or row.get("Wahlberechtigte insgesamt"))
    waehler = _num(row.get("Waehler") or row.get("Wähler") or row.get("Wählende"))
    ist = _num(row.get("AusWbez") or row.get("Ist.Wahlbezirke"))
    soll = _num(row.get("AnzWbez") or row.get("Soll.Wahlbezirke"))
    if wber <= 0 or waehler <= 0:
        for k, v in row.items():
            lk = str(k).lower().replace("ü", "ue").replace("ä", "ae")
            if wber <= 0 and "wahlberecht" in lk and "p" not in lk[-2:]:
                wber = _num(v)
            elif waehler <= 0 and ("waehler" in lk or "waehlende" in lk) and not lk.endswith("p"):
                waehler = _num(v)
    return {
        "counts": counts,
        "gueltig": gueltig,
        "wber": wber,
        "waehler": waehler,
        "ist": ist,
        "soll": soll,
        "raw": row,
    }


def _wid_from_afs_row(row: dict, awk_to_wkr: dict[str, str], addr_to_wkr: dict[str, str]) -> str | None:
    nummer = str(row.get("Nummer") or "").strip()
    if nummer in awk_to_wkr:
        return awk_to_wkr[nummer]
    if re.fullmatch(r"\d{4}", nummer) and nummer in awk_to_wkr:
        return awk_to_wkr[nummer]
    addr = str(row.get("Adresse") or "").strip()
    if addr in addr_to_wkr:
        return addr_to_wkr[addr]
    bez = str(row.get("Bezirk") or row.get("Bezirksnummer") or "").strip()
    agh = row.get("AghWkr") or row.get("Abgeordnetenhauswahlkreis")
    if not bez:
        m = re.match(r"^(\d{2})", addr)
        if m:
            bez = m.group(1)
    key = awk_code(bez, agh) if agh not in (None, "") else None
    if key and key in awk_to_wkr:
        return awk_to_wkr[key]
    # Adresse like AI0101 (aggregate WK stub)
    m = re.search(r"(\d{4})$", addr)
    if m and m.group(1) in awk_to_wkr:
        return awk_to_wkr[m.group(1)]
    return None


def parse_afs_aggregate(path: Path, addr_to_wkr: dict[str, str] | None = None) -> dict:
    """Parse AfS _A_ CSV (Land / Bezirk / AGH-WK). P01… or named parties.

    Always keep AnzWbez/AusWbez even when vote columns are still 0 — that is
    the Wahlbezirk Soll, not the 78 Wahlkreise.
    """
    rows, cols = read_csv_rows(path)
    if not rows:
        return {"land": {}, "wkr": {}, "bezirk": {}, "clock": None, "kind": "empty"}
    party_cols = map_party_columns(cols, AFS_2026_PXX)
    awk_to_wkr = load_awk_to_wkr()
    addr_to_wkr = addr_to_wkr or {}
    land: dict = {}
    wkrs: dict[str, dict] = {}
    bez: dict[str, dict] = {}
    clock = None

    for row in rows:
        art = str(row.get("Gebietsart") or "").strip()
        art_l = art.lower()
        clock = clock or _afs_clock(row)
        if "bundestag" in art_l or art_l in ("ost/west", "ostwest"):
            continue
        unit = _afs_unit(row, party_cols)
        if art_l == "bundesland" or str(row.get("Gebietsname") or "").strip().lower() == "berlin":
            if art_l == "bundesland" or not land:
                land = unit
            continue
        if art_l == "bezirk":
            code = str(row.get("Nummer") or "").strip().zfill(2)
            if code == "00":
                continue
            if not re.fullmatch(r"\d{2}", code):
                code = _bezirk_of_name(str(row.get("Gebietsname") or row.get("Bezirksname") or "")) or ""
            if code:
                bez[code] = unit
            continue
        if "abgeordnetenhauswahlkreis" in art_l or art_l in ("wahlkreis", "awk"):
            wid = _wid_from_afs_row(row, awk_to_wkr, addr_to_wkr)
            if wid:
                wkrs[wid] = unit
            continue
        # Some _A_ dumps omit Gebietsart; infer from Nummer/name.
        name = str(row.get("Gebietsname") or row.get("Name") or "")
        if _bezirk_of_name(name) and not re.search(r"\d", name[-3:]):
            code = _bezirk_of_name(name)
            if code:
                bez[code] = unit
            continue
        wid = _wid_from_afs_row(row, awk_to_wkr, addr_to_wkr)
        if wid:
            wkrs[wid] = unit
    has_votes = _num((land or {}).get("gueltig")) > 0 or any(
        _num(u.get("gueltig")) > 0 for u in wkrs.values()
    )
    if not has_votes:
        kind = "empty"
    elif wkrs:
        kind = "Z"
    else:
        kind = "A"
    return {"land": land, "wkr": wkrs, "bezirk": bez, "clock": clock, "kind": kind}


def parse_afs_precincts(path: Path, addr_to_wkr: dict[str, str] | None = None) -> dict[str, dict]:
    """Parse AfS _W_ CSV keyed by Adresse. Empty / uncounted → {}."""
    rows, cols = read_csv_rows(path)
    if not rows or _looks_empty_afs(rows):
        return {}
    party_cols = map_party_columns(cols, AFS_2026_PXX)
    awk_to_wkr = load_awk_to_wkr()
    addr_to_wkr = addr_to_wkr or {}
    out: dict[str, dict] = {}
    for row in rows:
        addr = str(row.get("Adresse") or row.get("Wahlbezirk") or "").strip()
        if not addr:
            continue
        art = str(row.get("Gebietsart") or "").lower()
        if art and "wahlbezirk" not in art and "urne" not in art and "brief" not in art:
            # _A_ file passed in by mistake
            if "wahlkreis" in art or "bezirk" in art or "bundesland" in art:
                continue
        counts, gueltig = counts_from_row(row, party_cols, ["Gueltig", "Gültige Stimmen"])
        if gueltig <= 0:
            continue
        wid = _wid_from_afs_row(row, awk_to_wkr, addr_to_wkr)
        wb_art = str(row.get("WBezArt") or row.get("Wahlbezirksart") or "").strip().upper()[:1]
        if not wb_art:
            m = re.search(r"\d{2}([WBU])", addr, re.I)
            wb_art = (m.group(1) if m else "W").upper()
            if wb_art == "U":
                wb_art = "W"
        out[addr] = {
            "id": addr,
            "wkr": wid,
            "counts": counts,
            "gueltig": gueltig,
            "art": wb_art or "W",
        }
    return out


def precincts_to_wkr(wb: dict[str, dict], panel: dict[str, dict]) -> dict[str, dict]:
    """Roll reported WBs up to WK live rows (ist = n reported)."""
    by: dict[str, dict] = {}
    for rec in wb.values():
        wid = rec.get("wkr")
        if not wid or wid not in panel:
            continue
        slot = by.setdefault(
            wid,
            {
                "counts": {p: 0.0 for p in PARTIES},
                "gueltig": 0.0,
                "ist": 0.0,
                "soll": 0.0,
                "wber": 0.0,
                "waehler": 0.0,
            },
        )
        for p in PARTIES:
            slot["counts"][p] += float(rec["counts"].get(p, 0.0))
        slot["gueltig"] += rec["gueltig"]
        slot["ist"] += 1
    # soll unknown at WB level — leave 0 so frac comes from volume
    return by


def allocate_from_nowcast(frac: dict[str, float], dirs: dict[str, int]) -> dict:
    votes = {p: float(frac.get(p, 0.0)) for p in MAIN}
    directs = {p: int(dirs.get(p, 0)) for p in MAIN}
    return allocate_be(votes, directs, bezirk_parties={"cdu", "spd", "linke"})


def build_step(
    panel,
    prior,
    live,
    land_prior,
    rng,
    *,
    bsw_direkt,
    erst_prior,
    prior_unc,
    wk_unc,
    state_draws,
    external,
    wb_live,
) -> dict:
    live_wkr = dict(live.get("wkr") or {})
    source = "wkr"
    if wb_live:
        rolled = precincts_to_wkr(wb_live, panel)
        if rolled:
            live_wkr = rolled
            source = "wahlbezirke"
    nc, diag = nowcast_wkrs(
        panel,
        prior,
        live_wkr,
        bsw_direkt=bsw_direkt,
        erst_prior=erst_prior,
        prior_unc_pp=prior_unc,
    )
    total_g = sum(panel[w]["gueltig_l1"] for w in panel) + EPS
    nc_land = {
        p: sum(panel[w]["gueltig_l1"] * nc[w]["nowcast"][p] for w in panel) / total_g
        for p in PARTIES
    }
    land_row = live.get("land") or {}
    frac_v = float(diag.get("frac_votes") or 0.0)
    if _num(land_row.get("gueltig")) > 0 and diag["n_wkr_touch"] == 0 and not wb_live:
        nc_land = _shares(land_row.get("counts") or nc_land)
        source = "land"
        soll_l = _num(land_row.get("soll"))
        ist_l = _num(land_row.get("ist"))
        if soll_l > 0:
            frac_v = max(frac_v, min(0.99, ist_l / soll_l))
        else:
            frac_v = max(frac_v, 0.15)
    unc = dict(diag["uncertainty"])
    nc_land, unc, pred = blend_with_external(nc_land, unc, frac_v, external, prior_unc)

    n_rep = sum(1 for w in nc.values() if w["reported"])
    n_tot = len(panel)
    extra_ist = len(wb_live) if wb_live else 0
    extra_soll = len(wb_live) if wb_live else 0
    ist, soll = precinct_ist_soll(
        land_row,
        nc,
        extra_ist=extra_ist,
        extra_soll=extra_soll,
        expected_soll=EXPECTED_PRECINCTS,
        n_wkr=n_tot,
    )
    frac_wb = (ist / soll) if soll else frac_v

    wber = _num(land_row.get("wber"))
    waehler = _num(land_row.get("waehler"))
    naive = 100.0 * waehler / wber if wber > 1000 and waehler > 0 else None
    n_complete = sum(1 for w in nc.values() if w.get("frac", 0) >= 0.999)
    brief_visible = any(
        (wb_live or {}).get(a, {}).get("art") == "B" for a in (wb_live or {})
    )
    turnout = turnout_mixture(
        prior_pp=TURNOUT_PRIOR,
        last_office_pp=LAST_OFFICE_TURNOUT,
        recent_federal_pp=RECENT_FEDERAL_TURNOUT,
        naive=naive,
        reported_frac=frac_wb,
        n_complete=n_complete,
        brief_visible=brief_visible,
        external_turnout=(external or {}).get("turnout"),
    )
    races, wkr_out = wkr_races(panel, nc, wk_unc, prior_unc)
    by_bez: dict[str, dict] = {}
    for wid, row in panel.items():
        bez = row.get("bezirk") or "?"
        slot = by_bez.setdefault(bez, {p: 0.0 for p in PARTIES} | {"g": 0.0})
        g = row["gueltig_l1"]
        slot["g"] += g
        for p in PARTIES:
            slot[p] += g * nc[wid]["nowcast"][p]
    bez_out = {}
    for bez, slot in by_bez.items():
        g = slot.pop("g") or EPS
        bez_out[bez] = {
            "nowcast": _pct({p: slot[p] / g for p in PARTIES}),
            "label": BEZ_NAMES.get(bez, bez),
        }
    entry_mc = night_entry_mc(
        _pct(nc_land),
        unc,
        races,
        rng,
        allocate=allocate_from_nowcast,
        base_seats=130,
        state_draws=state_draws,
        prior_unc_pp=prior_unc,
        by_bez_pct={b: row["nowcast"] for b, row in bez_out.items()},
        wkr_bez={wid: str(row.get("bezirk") or "") for wid, row in panel.items()},
        bezirk_parties=BE_BEZIRKSLISTE_PARTIES,
    )
    scen = night_scenario_probs(
        _pct(nc_land),
        unc,
        rng,
        state="BE",
        state_draws=state_draws,
        prior_unc_pp=prior_unc,
    )
    clock = live.get("clock") or datetime.now().astimezone().strftime("%Y-%m-%d %H:%M")
    kind = live.get("kind") or "empty"
    return {
        "frac_reported": round(frac_wb if soll else frac_v, 4),
        "n_reported": int(ist),
        "n_total": int(soll) if soll else EXPECTED_PRECINCTS,
        "clock": clock,
        "clock_source": "afs",
        "nowcast": _pct(nc_land),
        "naive": diag["naive"],
        "prior": _pct(land_prior),
        "truth": None,
        "mae_nowcast": None,
        "mae_naive": None,
        "learn_weight": diag["learn_weight"],
        "mix_live": diag.get("mix_live"),
        "mix_prior": diag.get("mix_prior"),
        "representativeness": None,
        "surprise": diag["surprise"],
        "nowcast_source": source,
        "uncertainty": unc,
        "turnout": turnout,
        "by_wkr": wkr_out,
        "by_bezirk": bez_out,
        "entry_mc": entry_mc,
        "scenario_probs": scen,
        "eval": None,
        "result_kind": kind,
        "n_wkr_touch": diag["n_wkr_touch"],
        "n_wkr_reported": int(n_rep),
        "n_wkr_total": int(n_tot),
        "prediction": pred,
        "uncertainty_note": {
            "phase": unc_phase(frac_wb if soll else frac_v),
            "land": (
                "Landes-± startet mit der Vorwahlprognose bzw. 18-Uhr-Prognose/"
                "Hochrechnung und schrumpft mit dem offenen Stimmenanteil."
            ),
            "wkr": "Wahlkreis-± = Regressionsband × offener Anteil in diesem Kreis.",
        },
    }


def run(prev_path: Path | None, external_path: Path | None) -> dict:
    panel = load_wkr_panel(PANEL_PATH)
    attach_bezirke(panel)
    alt, addr_to_wkr = load_btw2025_wk_shares(panel)
    if alt:
        blend_wk_history(panel, alt, weight_alt=0.55)
    else:
        for row in panel.values():
            row["shares_hist"] = dict(row["shares_l1"])
            row["hist_source"] = "agh2023"
    state_fc = _load_json(REPO / "output" / "forecast_state_be.json", FORECAST_STATE_URL)
    dist_fc = _load_json(REPO / "output" / "forecast_districts_be.json", FORECAST_DISTRICT_URL)
    draws_doc = _load_json_optional(
        REPO / "output" / "forecast_state_be_draws.json", FORECAST_DRAWS_URL
    )
    state_draws = load_state_draws(draws_doc)
    land_prior = statewide_prior(state_fc)
    if state_draws is not None:
        m = state_draws.mean(axis=0)
        land_prior = _shares({p: float(m[i]) for i, p in enumerate(PARTIES)})
    prior_unc = prior_uncertainty_pp(state_fc)
    prior, erst_prior, bsw_direkt = district_priors(
        dist_fc, panel, land_prior, hist_key="shares_hist"
    )
    wk_unc = district_wk_unc_pp(dist_fc)
    zweit_a = LIVE_DIR / "Datenexport_AGH2026_Zweitstimme_A_BE.csv"
    erst_a = LIVE_DIR / "Datenexport_AGH2026_Erststimme_A_BE.csv"
    zweit_w = LIVE_DIR / "Datenexport_AGH2026_Zweitstimme_W_BE.csv"
    erst_w = LIVE_DIR / "Datenexport_AGH2026_Erststimme_W_BE.csv"
    live = parse_afs_aggregate(zweit_a, addr_to_wkr)
    erst_agg = parse_afs_aggregate(erst_a, addr_to_wkr)
    for wid, unit in (erst_agg.get("wkr") or {}).items():
        slot = live.setdefault("wkr", {}).setdefault(wid, {})
        slot["erst_counts"] = unit.get("counts")
        slot["erst_gueltig"] = unit.get("gueltig")
    wb = parse_afs_precincts(zweit_w, addr_to_wkr)
    wb_e = parse_afs_precincts(erst_w, addr_to_wkr)
    ankunft = load_ankunft(LIVE_DIR)
    for addr, rec in wb_e.items():
        if addr in wb:
            wb[addr]["erst_counts"] = rec["counts"]
            wb[addr]["erst_gueltig"] = rec["gueltig"]
    external = load_external("be", external_path or EXTERNAL_JSON)
    rng = np_rng()
    step = build_step(
        panel,
        prior,
        live,
        land_prior,
        rng,
        bsw_direkt=bsw_direkt,
        erst_prior=erst_prior,
        prior_unc=prior_unc,
        wk_unc=wk_unc,
        state_draws=state_draws,
        external=external,
        wb_live=wb,
    )
    prev = None
    if prev_path and prev_path.exists():
        try:
            prev = json.loads(prev_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            prev = None
    steps = merge_history(union_history_payloads(prev), step)
    hist_note = (
        f"WK-Historie: AGH2023"
        + (f" + BTW2025 ({len(alt)} WK)" if alt else " (BTW2025-Remap nicht geladen)")
    )
    payload = {
        "election": "AGH2026",
        "state": "be",
        "state_label": "Berlin",
        "last_election": {
            "year": 2023,
            "label": "AGH 2023",
            "turnout": LAST_OFFICE_TURNOUT,
            "parliament_size": 159,
        },
        "baseline": (
            f"π₀ = Landesprognose; {hist_note}; Live = AfS _A_"
            + ("/_W_" if wb else " + Ankunftstafel (keine _W_-Stimmen)")
        ),
        "parties": list(PARTIES),
        "party_labels": PARTY_LABELS,
        "n_precincts": int(step["n_total"] or EXPECTED_PRECINCTS),
        "n_wkr": len(panel),
        "match": {
            "n_truth": len(panel),
            "n_l1": len(panel),
            "n_matched": len(alt) if alt else len(panel),
            "match_rate": round((len(alt) / len(panel)) if alt else 1.0, 4),
        },
        "generated_at": github_now(),
        "prior_uncertainty_pp": prior_unc,
        "model": {
            "description": (
                "Berlin AGH 2026 Live-Nowcast. Lokale Priors: AGH 2023 und "
                "BTW 2025 auf 2026-Wahlkreisen, ausgerichtet auf die "
                "zweitstimme.org-Landesprognose. Live-Feed AfS _A_ "
                "(Land/Bezirk/WK) plus Ankunftstafel der Präsentation "
                "(eingegangene Wahlbezirke ohne Stimmen). _W_-CSV erst mit "
                "dem vorläufigen Ergebnis. 18-Uhr-Prognose/Hochrechnung aus "
                "wahlabend_external.json ersetzt π₀ und Szenarien, dann "
                "Mischung mit der Auszählung."
            )
        },
        "call_threshold": 0.90,
        "hard_call_threshold": 0.999,
        "geo_units": {
            "land": [{"id": "BE", "label": "Berlin"}],
            "bezirk": [
                {"id": k, "label": v} for k, v in BEZ_NAMES.items()
            ],
            "wkr": [
                {
                    "id": wid,
                    "label": f"{wid} {panel[wid]['name']}",
                    "bezirk": panel[wid].get("bezirk"),
                }
                for wid in sorted(panel, key=lambda x: int(x) if x.isdigit() else 99)
            ],
        },
        "scenarios": {
            "live": {
                "label": "AfS Live-CSV",
                "steps": steps,
                "wkr_calls": {},
                "summary": {},
            }
        },
        "scenario": "live",
        "features": {"bezirkslisten": True, "listen_einzug": True},
        "listen_mode": "bezirk",
        "listen_roster_2026": load_listen_roster(
            REPO / "berlin" / "candidates" / "listenkandidaten_2026.csv"
        ),
        "listen_roster_note": (
            "Berlin 2026: CDU/SPD/Linke mit Bezirkslisten, übrige Landesliste. "
            "Einzug = 5 % oder ein Direktmandat."
        ),
        "direkt_candidates_2026": load_direkt_roster(
            REPO / "berlin" / "candidates" / "direktkandidaten_2026.csv"
        ),
        "live": {
            **live_meta(
                step.get("result_kind") or "empty",
                KIND_LABEL.get(step.get("result_kind") or "empty", "AfS"),
                step,
                f"{AFS_BASES[0]}/Datenexport_AGH2026_Zweitstimme_A_BE.csv",
            ),
            "n_wb_reported": len(wb),
            "wb_level": "wahlbezirke" if wb else "wkr",
            "wb_level_note": None if wb else BE_WB_NIGHT_NOTE,
            "n_ankunft": int(ankunft.get("n") or 0),
            "ankunft_source": AFS_INDEX,
            "ankunft_latest": ankunft.get("latest") or [],
        },
        "precincts": live_precincts(panel, live.get("wkr") or {})
        + [
            {
                "id": a,
                "wkr": rec.get("wkr"),
                "art": rec.get("art") or "W",
                "name": a,
                "bezirk": None,
                "gueltig": int(rec.get("gueltig") or 0),
                "counts": {p: int(round(float(rec["counts"].get(p, 0)))) for p in PARTIES},
            }
            for a, rec in list(wb.items())[:400]
        ],
    }
    return payload


def np_rng():
    import numpy as np

    return np.random.default_rng(20260920)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=OUT_DIR / "wahlabend_nowcast_be_live.json")
    ap.add_argument("--prev", type=Path, default=None)
    ap.add_argument("--external", type=Path, default=None)
    args = ap.parse_args()
    LIVE_DIR.mkdir(parents=True, exist_ok=True)
    prev = args.prev or args.out
    payload = run(prev, args.external)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    live = payload["live"]
    step = payload["scenarios"]["live"]["steps"][-1]
    print(
        f"wrote {args.out}  kind={live.get('result_kind')}  "
        f"{live.get('ist_wb')}/{live.get('soll_wb')}  "
        f"CDU={step['nowcast'].get('cdu')}  "
        f"src={live.get('prediction_source_label')}  "
        f"steps={len(payload['scenarios']['live']['steps'])}"
    )


if __name__ == "__main__":
    main()
